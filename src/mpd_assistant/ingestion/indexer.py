"""Idempotent knowledge indexing into document_chunks. Re-running only re-embeds documents that changed.

Text documents are read directly; scanned PDFs and images go through OCR when an engine is configured, and their
chunks keep the OCR confidence. Files removed from the knowledge folder are removed from the index.
"""

import asyncio
import hashlib
import io
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional

from sqlalchemy import Engine, delete, select
from sqlalchemy.orm import Session as DbSession

from mpd_assistant.core.models import EMBEDDING_DIM, DocumentChunk
from mpd_assistant.ingestion.chunker import chunk_text
from mpd_assistant.multimodal.ocr import OcrEngine
from mpd_assistant.observability.metrics import INGEST_DOCUMENTS
from mpd_assistant.retrieval.embeddings import Embedder

TEXT_SUFFIXES = {".txt", ".md"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff"}
SUPPORTED = TEXT_SUFFIXES | IMAGE_SUFFIXES | {".docx", ".pdf"}
MAX_PDF_PAGES = 200


class IndexError_(RuntimeError):
    pass


@dataclass
class Extracted:
    text: str
    ocr_confidence: Optional[float] = None


def read_document(path: Path) -> str:
    """Text-layer extraction only (no OCR). Raises IndexError_ for unsupported types."""
    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".docx":
        from docx import Document

        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            parts.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
        return "\n\n".join(p for p in parts if p.strip())
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages[:MAX_PDF_PAGES])
    if suffix in IMAGE_SUFFIXES:
        return ""
    raise IndexError_(f"Unsupported file type: {suffix}")


def _pdf_page_images(path: Path) -> List[bytes]:
    from pypdf import PdfReader

    images = []
    for page in PdfReader(str(path)).pages[:MAX_PDF_PAGES]:
        largest = max(page.images, key=lambda img: len(img.data), default=None)
        if largest is not None:
            images.append(largest.data)
    return images


def _to_png(data: bytes) -> bytes:
    from PIL import Image

    with Image.open(io.BytesIO(data)) as img:
        out = io.BytesIO()
        img.convert("RGB").save(out, format="PNG")
        return out.getvalue()


async def extract(path: Path, ocr: Optional[OcrEngine] = None) -> Extracted:
    text = await asyncio.to_thread(read_document, path)
    if text.strip():
        return Extracted(text)
    suffix = path.suffix.lower()
    if ocr is None or suffix not in IMAGE_SUFFIXES | {".pdf"}:
        return Extracted("")
    images = await asyncio.to_thread(_pdf_page_images, path) if suffix == ".pdf" else [path.read_bytes()]
    texts: List[str] = []
    confidences: List[float] = []
    for raw in images:
        result = await ocr.read(await asyncio.to_thread(_to_png, raw))
        texts.append(result.text)
        if result.confidence is not None:
            confidences.append(result.confidence)
    return Extracted("\n\n".join(texts), min(confidences) if confidences else None)


async def index_document(
    engine: Engine,
    tenant_id: str,
    source: str,
    text: str,
    embedder: Embedder,
    expected_dim: int = EMBEDDING_DIM,
    ocr_confidence: Optional[float] = None,
) -> str:
    """Returns 'indexed', 'unchanged' or 'empty'."""
    doc_hash = hashlib.sha256(text.encode("utf-8")).hexdigest()
    with DbSession(engine) as db:
        existing = db.scalar(
            select(DocumentChunk.doc_hash).where(DocumentChunk.tenant_id == tenant_id, DocumentChunk.source == source).limit(1)
        )
    if existing == doc_hash:
        return "unchanged"

    chunks = chunk_text(text)
    if not chunks:
        return "empty"
    vectors = await embedder.embed(chunks)
    if any(len(v) != expected_dim for v in vectors):
        raise IndexError_(f"Embedding dimension mismatch (expected {expected_dim})")

    with DbSession(engine) as db:
        db.execute(delete(DocumentChunk).where(DocumentChunk.tenant_id == tenant_id, DocumentChunk.source == source))
        for i, (chunk, vector) in enumerate(zip(chunks, vectors, strict=True)):
            db.add(
                DocumentChunk(
                    tenant_id=tenant_id, source=source, doc_hash=doc_hash, chunk_index=i, text=chunk, embedding=vector, ocr_confidence=ocr_confidence
                )
            )
        db.commit()
    return "indexed"


def remove_document(engine: Engine, tenant_id: str, source: str) -> int:
    with DbSession(engine) as db:
        removed = db.execute(delete(DocumentChunk).where(DocumentChunk.tenant_id == tenant_id, DocumentChunk.source == source)).rowcount or 0
        db.commit()
        return removed


def indexed_sources(engine: Engine, tenant_id: str) -> List[str]:
    with DbSession(engine) as db:
        return list(db.scalars(select(DocumentChunk.source).where(DocumentChunk.tenant_id == tenant_id).distinct()))


async def index_directory(
    engine: Engine,
    tenant_id: str,
    directory: Path,
    embedder: Embedder,
    expected_dim: int = EMBEDDING_DIM,
    ocr: Optional[OcrEngine] = None,
    prune: bool = True,
) -> Dict[str, str]:
    """Index every supported file; with prune, drop index entries for files that no longer exist."""
    results: Dict[str, str] = {}
    present = set()
    for path in sorted(directory.iterdir()):
        if path.name.startswith("~$") or path.suffix.lower() not in SUPPORTED:
            continue
        present.add(path.name)
        if ocr is None and path.suffix.lower() in IMAGE_SUFFIXES:
            continue  # nothing to read without OCR; keep any entry indexed earlier
        try:
            extracted = await extract(path, ocr)
            status = await index_document(engine, tenant_id, path.name, extracted.text, embedder, expected_dim, extracted.ocr_confidence)
        except Exception as exc:
            status = f"failed:{type(exc).__name__}"
        results[path.name] = status
        INGEST_DOCUMENTS.labels(status.split(":")[0]).inc()
    if prune:
        for source in indexed_sources(engine, tenant_id):
            if source not in present:
                remove_document(engine, tenant_id, source)
                results[source] = "removed"
                INGEST_DOCUMENTS.labels("removed").inc()
    return results
