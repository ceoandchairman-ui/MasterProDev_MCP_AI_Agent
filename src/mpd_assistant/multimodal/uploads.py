"""File uploads used as per-conversation context.

Files are checked by content (magic bytes), not by name. Extracted text is pseudonymized with the session vault,
stored for the conversation only (TTL), and handed to agents inside an untrusted block, so instructions in a file
can never trigger a tool or be written to memory.
"""

import asyncio
import io
import uuid
from dataclasses import dataclass
from typing import List

from mpd_assistant.core.kv import KeyValueStore
from mpd_assistant.multimodal.ocr import OcrEngine, OcrError, VisionModelOcr, strip_metadata
from mpd_assistant.privacy.vault import PiiVault
from mpd_assistant.security.untrusted import looks_like_injection, wrap

MAX_BYTES = 10 * 1024 * 1024
MAX_FILES_PER_THREAD = 5
MAX_TEXT_PER_FILE = 12000
MAX_CONTEXT_CHARS = 6000
TTL_SECONDS = 24 * 3600
MIN_OCR_CONFIDENCE = 0.5


class UploadRejected(ValueError):
    def __init__(self, reason: str):
        super().__init__(reason)
        self.reason = reason


@dataclass
class UploadResult:
    id: str
    kind: str
    chars: int
    note: str = ""  # e.g. "low_ocr_confidence", "possible_instructions"


def sniff(data: bytes) -> str | None:
    if data.startswith(b"%PDF-"):
        return "pdf"
    if data.startswith(b"PK\x03\x04") and b"word/" in data[:4000]:
        return "docx"
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image"
    if data.startswith(b"\xff\xd8\xff"):
        return "image"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image"
    if b"\x00" not in data[:4096]:
        try:
            data[:4096].decode("utf-8")
            return "text"
        except UnicodeDecodeError:
            return None
    return None


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise UploadRejected("encrypted_pdf")
    return "\n\n".join((page.extract_text() or "") for page in reader.pages[:50])


def _pdf_images(data: bytes, max_pages: int = 10) -> List[bytes]:
    from pypdf import PdfReader

    images: List[bytes] = []
    for page in PdfReader(io.BytesIO(data)).pages[:max_pages]:
        largest = max(page.images, key=lambda img: len(img.data), default=None)
        if largest is not None:
            images.append(largest.data)
    return images


def _docx_text(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        parts.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
    return "\n".join(p for p in parts if p.strip())


class UploadService:
    def __init__(self, kv: KeyValueStore, vault: PiiVault, ocr: OcrEngine | None = None, max_bytes: int = MAX_BYTES):
        self._kv, self._vault, self._ocr, self._max = kv, vault, ocr, max_bytes

    @property
    def max_bytes(self) -> int:
        return self._max

    def _key(self, thread_id: str) -> str:
        return f"uploads:{thread_id}"

    async def _image_text(self, data: bytes) -> tuple[str, str]:
        if self._ocr is None:
            raise UploadRejected("images_not_supported")
        try:
            clean = await asyncio.to_thread(strip_metadata, data)
        except Exception as exc:
            raise UploadRejected("unreadable_image") from exc
        try:
            result = await self._ocr.read(clean)
            if len(result.text.strip()) < 20 and isinstance(self._ocr, VisionModelOcr):
                return await self._ocr.describe(clean), ""  # not a document: describe it instead
        except OcrError as exc:
            raise UploadRejected("unreadable_image") from exc
        note = "low_ocr_confidence" if result.confidence is not None and result.confidence < MIN_OCR_CONFIDENCE else ""
        return result.text, note

    async def _scanned_pdf_text(self, data: bytes) -> tuple[str, str]:
        """A scanned PDF is usually one image per page: OCR the embedded page images."""
        if self._ocr is None:
            raise UploadRejected("scanned_pdf_not_supported")
        images = await asyncio.to_thread(_pdf_images, data)
        if not images:
            raise UploadRejected("no_text_found")
        texts, notes = [], set()
        for image in images:
            text, note = await self._image_text(image)
            texts.append(text)
            if note:
                notes.add(note)
        return "\n\n".join(texts), ",".join(sorted(notes))

    async def accept(self, thread_id: str, session_id: str, filename: str, data: bytes) -> UploadResult:
        if not data:
            raise UploadRejected("empty_file")
        if len(data) > self._max:
            raise UploadRejected("too_large")
        stored = self._kv.get(self._key(thread_id)) or []
        if len(stored) >= MAX_FILES_PER_THREAD:
            raise UploadRejected("too_many_files")
        kind = sniff(data)
        if kind is None:
            raise UploadRejected("unsupported_type")

        note = ""
        try:
            if kind == "pdf":
                text = await asyncio.to_thread(_pdf_text, data)
                if not text.strip():
                    text, note = await self._scanned_pdf_text(data)
            elif kind == "docx":
                text = await asyncio.to_thread(_docx_text, data)
            elif kind == "image":
                text, note = await self._image_text(data)
            else:
                text = data.decode("utf-8", errors="replace")
        except UploadRejected:
            raise
        except Exception as exc:
            raise UploadRejected("unreadable_file") from exc

        text = text.strip()[:MAX_TEXT_PER_FILE]
        if not text:
            raise UploadRejected("no_text_found")
        safe = self._vault.tokenize(session_id, text)
        if looks_like_injection(safe):
            note = note or "possible_instructions"
        upload_id = uuid.uuid4().hex[:12]
        safe_name = "".join(c for c in (filename or "file")[-80:] if c.isalnum() or c in "._- ") or "file"
        stored.append({"id": upload_id, "name": safe_name, "kind": kind, "text": safe})
        self._kv.set(self._key(thread_id), stored, TTL_SECONDS)
        return UploadResult(upload_id, kind, len(safe), note)

    def context_for(self, thread_id: str) -> str:
        """Most recent files first, wrapped as untrusted data, within a character budget."""
        blocks: List[str] = []
        used = 0
        for item in reversed(self._kv.get(self._key(thread_id)) or []):
            remaining = MAX_CONTEXT_CHARS - used
            if remaining <= 200:
                break
            body = item["text"][:remaining]
            blocks.append(wrap(body, f"upload:{item['name']}"))
            used += len(body)
        return "\n".join(blocks)

    def remove(self, thread_id: str, upload_id: str | None = None) -> int:
        stored = self._kv.get(self._key(thread_id)) or []
        kept = [i for i in stored if upload_id is not None and i["id"] != upload_id]
        self._kv.set(self._key(thread_id), kept, TTL_SECONDS)
        return len(stored) - len(kept)


__all__ = ["UploadRejected", "UploadResult", "UploadService", "sniff"]
